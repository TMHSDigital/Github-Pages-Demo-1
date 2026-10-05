"""
Builds assets/downloads/prime-cost-tracker.xlsx: a month-by-month prime cost tracker that
uses the same formulas and rule-of-thumb bands as the calculators on the site.

Run: pip install -r scripts/requirements.txt && python scripts/make-tracker-template.py
The bands are TODO(verify) like the ones in index.html (#27): change them on the Settings
sheet and in js/calc-math.js together.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / 'assets' / 'downloads' / 'prime-cost-tracker.xlsx'
ROOT = Path(__file__).resolve().parent.parent
# The site's address, from the tool registry
SITE = json.loads((ROOT / 'tools' / 'tools.json').read_text(encoding='utf-8'))['site'] + 'tools/prime-cost-calculator/'

FONT = 'Arial'
NAVY = '203164'
INPUT_FILL = PatternFill('solid', fgColor='FFFF00')
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
EXAMPLE_FILL = PatternFill('solid', fgColor='EEF1F7')
thin = Side(style='thin', color='C9D3E6')
GRID = Border(bottom=thin)

f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_text = Font(name=FONT, size=10, color='333333')
f_head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_input = Font(name=FONT, size=10, color='0000FF')
f_calc = Font(name=FONT, size=10, color='000000')
f_bold = Font(name=FONT, size=10, bold=True)

USD = '$#,##0;($#,##0);-'
PCT = '0.0%;(0.0%);-'
PTS = '+0.0;-0.0;0.0'

wb = Workbook()

# Settings: the bands, each in its own labelled cell
st = wb.active
st.title = 'Settings'
st['A1'] = 'Bands for prime cost % (food and beverage + labor, as a share of sales)'
st['A1'].font = f_bold
rows = [
    ('Strong: prime cost below', 0.60),
    ('On target: below', 0.65),
    ('Watch closely: below', 0.70),
]
for i, (label, value) in enumerate(rows, start=3):
    st.cell(i, 1, label).font = f_text
    c = st.cell(i, 2, value)
    c.font, c.fill, c.number_format = f_input, INPUT_FILL, '0%'
st['A6'] = 'At or above the last band: Needs attention.'
st['A6'].font = f_text
st['A8'] = ('Source: common rules of thumb, the same bands the calculators at ' + SITE + ' use. '
            'Healthy ranges differ by concept, so set your own; your own month-to-month trend says more than any benchmark.')
st['A8'].font = f_text
st['A8'].alignment = Alignment(wrap_text=True, vertical='top')
st.merge_cells('A8:D10')
st.column_dimensions['A'].width = 34
st.column_dimensions['B'].width = 12
for col in 'CD':
    st.column_dimensions[col].width = 30

# Tracker
ws = wb.create_sheet('Tracker', 0)
ws['A1'] = 'Prime cost tracker'
ws['A1'].font = f_title
ws['A2'] = ('Fill in the yellow cells (blue text) from each month\'s P&L; everything else calculates. '
            'Row 5 is an example to show the format: overwrite it or delete its numbers. Bands are set on the Settings sheet.')
ws['A2'].font = f_text
ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A2:L2')
ws.row_dimensions[2].height = 30

headers = [
    ('Month', 14), ('Sales ($)', 13), ('Food & beverage cost ($)', 15), ('Labor cost ($)', 13),
    ('Other controllable costs ($)', 15), ('Food & bev %', 11), ('Labor %', 10), ('Prime cost ($)', 13),
    ('Prime cost %', 11), ('Band', 15), ('Left after these costs ($)', 14), ('Change vs previous month (points)', 15),
]
HEAD = 4
for col, (name, width) in enumerate(headers, start=1):
    c = ws.cell(HEAD, col, name)
    c.font, c.fill = f_head, HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.column_dimensions[get_column_letter(col)].width = width
ws.row_dimensions[HEAD].height = 44

FIRST, LAST = 5, 5 + 24  # the example row plus two years of months
for r in range(FIRST, LAST + 1):
    for col in range(1, 6):  # inputs
        c = ws.cell(r, col)
        c.font, c.fill, c.border = f_input, INPUT_FILL, GRID
        c.number_format = '@' if col == 1 else USD
    has = f'$B{r}>0'
    formulas = {
        6: (f'=IF({has},C{r}/B{r},"")', PCT),
        7: (f'=IF({has},D{r}/B{r},"")', PCT),
        8: (f'=IF({has},C{r}+D{r},"")', USD),
        9: (f'=IF({has},H{r}/B{r},"")', PCT),
        10: (f'=IF(I{r}="","",IF(I{r}<Settings!$B$3,"Strong",IF(I{r}<Settings!$B$4,"On target",'
             f'IF(I{r}<Settings!$B$5,"Watch closely","Needs attention"))))', 'General'),
        11: (f'=IF({has},B{r}-C{r}-D{r}-E{r},"")', USD),
        12: ('' if r == FIRST else f'=IF(OR(I{r}="",I{r - 1}=""),"",(I{r}-I{r - 1})*100)', PTS),
    }
    for col, (formula, fmt) in formulas.items():
        c = ws.cell(r, col, formula or None)
        c.font, c.number_format, c.border = f_calc, fmt, GRID

# The example row: the calculator's worked example month
for col, value in enumerate(['Example', 100000, 31000, 32000, 12000], start=1):
    ws.cell(FIRST, col, value)
ws.cell(FIRST, 1).comment = Comment('Example values from the calculator\'s worked example. Replace with your own month, e.g. "Sep 2026".', 'TMHS')
for col in range(6, 13):
    ws.cell(FIRST, col).fill = EXAMPLE_FILL

ws.freeze_panes = ws.cell(FIRST, 2)
ws.cell(LAST + 2, 1, 'Prime cost % = (food and beverage cost + labor cost) / sales. '
                     'Change is in percentage points against the row above.').font = f_text

# openpyxl stores formulas without results; have Excel and Sheets compute them on open
wb.calculation.fullCalcOnLoad = True

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print('wrote', OUT)
